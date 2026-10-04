import re
import shutil
import tempfile
import zipfile
from pathlib import Path

from openpyxl import load_workbook


MAX_SUPPORTED_FONT_FAMILY = 14
REPLACEMENT_FONT_FAMILY = 2

FONT_FAMILY_PATTERN = re.compile(
    rb'(<(?:\w+:)?family[^>]*\bval=")(\d+)("[^>]*/?>)'
)


def _normalise_font_families(
    extracted_directory,
):
    """
    Normalise unsupported font-family metadata across
    XML components in an extracted XLSX package.

    Only values greater than openpyxl's supported maximum
    are changed.
    """

    adjustments = []

    for xml_path in extracted_directory.rglob(
        "*.xml"
    ):

        original_content = xml_path.read_bytes()

        file_adjustments = 0

        def replace_family(match):

            nonlocal file_adjustments

            family_value = int(
                match.group(2)
            )

            if (
                family_value
                <= MAX_SUPPORTED_FONT_FAMILY
            ):
                return match.group(0)

            file_adjustments += 1

            return (
                match.group(1)
                + str(
                    REPLACEMENT_FONT_FAMILY
                ).encode("ascii")
                + match.group(3)
            )

        updated_content = (
            FONT_FAMILY_PATTERN.sub(
                replace_family,
                original_content,
            )
        )

        if file_adjustments == 0:
            continue

        xml_path.write_bytes(
            updated_content
        )

        relative_path = (
            xml_path.relative_to(
                extracted_directory
            )
        )

        adjustments.append(
            {
                "component": str(
                    relative_path
                ).replace("\\", "/"),
                "adjustment": (
                    "Unsupported font-family "
                    "metadata normalised"
                ),
                "occurrences": file_adjustments,
            }
        )

    return adjustments


def create_compatible_audit_copy(
    source_path,
):
    """
    Create an internal audit copy of an XLSX workbook.

    The original workbook is never modified.

    If openpyxl cannot read the workbook because of
    supported compatibility issues, an internal copy is
    created with only the permitted metadata adjustments.
    """

    source_path = Path(
        source_path
    )

    temporary_directory = Path(
        tempfile.mkdtemp(
            prefix="excel_auditor_"
        )
    )

    copied_path = (
        temporary_directory
        / source_path.name
    )

    shutil.copy2(
        source_path,
        copied_path,
    )

    # First try the workbook without any adjustment.
    try:

        workbook = load_workbook(
            copied_path,
            read_only=True,
        )

        workbook.close()

        return {
            "path": copied_path,
            "adjustment_applied": False,
            "adjustments": [],
            "temporary_directory": (
                temporary_directory
            ),
        }

    except ValueError:

        pass

    extracted_directory = (
        temporary_directory
        / "extracted"
    )

    with zipfile.ZipFile(
        copied_path,
        "r",
    ) as archive:

        archive.extractall(
            extracted_directory
        )

    adjustments = (
        _normalise_font_families(
            extracted_directory
        )
    )

    if not adjustments:

        raise ValueError(
            "The workbook could not be made "
            "compatible using the supported "
            "formatting adjustments."
        )

    rebuilt_path = (
        temporary_directory
        / f"compatible_{source_path.name}"
    )

    with zipfile.ZipFile(
        rebuilt_path,
        "w",
        zipfile.ZIP_DEFLATED,
    ) as archive:

        for file_path in (
            extracted_directory.rglob("*")
        ):

            if not file_path.is_file():
                continue

            archive.write(
                file_path,
                file_path.relative_to(
                    extracted_directory
                ),
            )

    # Validation: the rebuilt workbook must now be
    # readable before it can ever be supplied to the
    # audit engine.
    workbook = load_workbook(
        rebuilt_path,
        read_only=True,
    )

    workbook.close()

    return {
        "path": rebuilt_path,
        "adjustment_applied": True,
        "adjustments": adjustments,
        "temporary_directory": (
            temporary_directory
        ),
    }