from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter


REPORT_SHEETS = [
    "Audit Summary",
    "Findings",
    "Finding Detail",
    "Audit Information",
]


NAVY = "17324D"
SLATE = "44546A"
LIGHT_BACKGROUND = "F3F6F8"
WHITE = "FFFFFF"

HIGH_PRIORITY_FILL = "F4CCCC"
REVIEW_FILL = "FCE5CD"
INFORMATION_FILL = "D9EAF7"


def _get_priority_sorted_findings(
    findings,
):
    """
    Return findings in review order without changing
    their original finding IDs.
    """

    priority_order = {
        "High Priority": 0,
        "Review Recommended": 1,
        "Information": 2,
    }

    return sorted(
        findings,
        key=lambda finding: (
            priority_order.get(
                finding.priority,
                99,
            ),
            finding.finding_id,
        ),
    )



def build_audit_report(audit_result):
    """
    Build an Excel audit report from an AuditResult.
    """

    workbook = Workbook()

    summary_sheet = workbook.active
    summary_sheet.title = "Audit Summary"

    findings_sheet = workbook.create_sheet(
        "Findings"
    )

    detail_sheet = workbook.create_sheet(
        "Finding Detail"
    )

    information_sheet = workbook.create_sheet(
        "Audit Information"
    )

    _build_summary_sheet(
        summary_sheet,
        audit_result,
    )

    _build_findings_sheet(
        findings_sheet,
        audit_result,
    )

    _build_detail_sheet(
        detail_sheet,
        audit_result,
    )

    _build_information_sheet(
        information_sheet,
        audit_result,
    )

    _autofit_and_wrap(
        findings_sheet,
        max_width=35,
    )

    _autofit_and_wrap(
        detail_sheet,
        max_width=45,
    )

    detail_sheet.column_dimensions["G"].width = 30
    detail_sheet.column_dimensions["H"].width = 42
    detail_sheet.column_dimensions["I"].width = 45
    detail_sheet.column_dimensions["K"].width = 45

    _adjust_detail_row_heights(
    detail_sheet
    )

  
    return workbook


def _build_summary_sheet(
    worksheet,
    audit_result,
):
    worksheet.sheet_view.showGridLines = False

    # Main title
    worksheet.merge_cells("A1:F2")
    worksheet["A1"] = "EXCEL WORKBOOK AUDIT REPORT"
    worksheet["A1"].font = Font(
        size=20,
        bold=True,
        color=WHITE,
    )
    worksheet["A1"].fill = PatternFill(
        "solid",
        fgColor=NAVY,
    )
    worksheet["A1"].alignment = Alignment(
        vertical="center",
    )

    # Workbook name
    worksheet.merge_cells("A3:F3")
    worksheet["A3"] = audit_result.workbook_name
    worksheet["A3"].font = Font(
        size=11,
        italic=True,
        color=SLATE,
    )

    # Section heading
    worksheet.merge_cells("A5:F5")
    worksheet["A5"] = "AUDIT OVERVIEW"
    worksheet["A5"].font = Font(
        size=12,
        bold=True,
        color=WHITE,
    )
    worksheet["A5"].fill = PatternFill(
        "solid",
        fgColor=SLATE,
    )

    # Metric cards
    metrics = [
        (
            "A7:B7",
            "A8:B8",
            audit_result.total_findings,
            "Total Findings",
            LIGHT_BACKGROUND,
        ),
        (
            "C7:D7",
            "C8:D8",
            audit_result.high_priority_count,
            "High Priority",
            HIGH_PRIORITY_FILL,
        ),
        (
            "E7:F7",
            "E8:F8",
            audit_result.review_recommended_count,
            "Review Recommended",
            REVIEW_FILL,
        ),
    ]

    for (
        value_range,
        label_range,
        value,
        label,
        fill_colour,
    ) in metrics:

        worksheet.merge_cells(value_range)
        worksheet.merge_cells(label_range)

        value_cell = worksheet[
            value_range.split(":")[0]
        ]
        label_cell = worksheet[
            label_range.split(":")[0]
        ]

        value_cell.value = value
        label_cell.value = label

        value_cell.font = Font(
            size=22,
            bold=True,
            color=NAVY,
        )

        label_cell.font = Font(
            size=10,
            bold=True,
            color=SLATE,
        )

        fill = PatternFill(
            "solid",
            fgColor=fill_colour,
        )

        value_cell.fill = fill
        label_cell.fill = fill

        value_cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        label_cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    # Secondary metrics
    thin_border = Side(
        style="thin",
        color="B7C3CE",
    )

    card_border = Border(
        left=thin_border,
        right=thin_border,
        top=thin_border,
        bottom=thin_border,
    )

    worksheet["A10"] = "Checks Completed"
    worksheet["B10"] = audit_result.checks_completed

    worksheet["D10"] = "Worksheets Affected"
    worksheet["E10"] = audit_result.worksheets_affected

    for coordinate in (
        "A10",
        "B10",
        "D10",
        "E10",
    ):
        worksheet[coordinate].border = card_border
        worksheet[coordinate].fill = PatternFill(
            "solid",
            fgColor=LIGHT_BACKGROUND,
        )
        worksheet[coordinate].alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for coordinate in ("A10", "D10"):
        worksheet[coordinate].font = Font(
            bold=True,
            color=SLATE,
        )

    for coordinate in ("B10", "E10"):
        worksheet[coordinate].font = Font(
            size=14,
            bold=True,
            color=NAVY,
        )

    worksheet.row_dimensions[10].height = 28

    # Review guidance
    worksheet.merge_cells("A12:F12")
    worksheet["A12"] = "REVIEW GUIDANCE"
    worksheet["A12"].font = Font(
        bold=True,
        color=WHITE,
    )
    worksheet["A12"].fill = PatternFill(
        "solid",
        fgColor=SLATE,
    )

    worksheet.merge_cells("A13:F14")

    if audit_result.high_priority_count > 0:
        guidance_text = (
            f"{audit_result.high_priority_count} high-priority "
            f"{'finding requires' if audit_result.high_priority_count == 1 else 'findings require'} "
            "attention. Review these first before relying on "
            "the workbook for analysis or decisions."
        )

    elif audit_result.review_recommended_count > 0:
        guidance_text = (
            "No high-priority issues were detected, but "
            f"{audit_result.review_recommended_count} "
            f"{'finding should' if audit_result.review_recommended_count == 1 else 'findings should'} "
            "be reviewed to confirm whether the identified "
            "patterns are intentional."
        )

    else:
        guidance_text = (
            "No issues were detected by the enabled "
            "audit checks."
        )

    worksheet["A13"] = guidance_text
    worksheet["A13"].font = Font(
        color=NAVY,
    )
    worksheet["A13"].fill = PatternFill(
        "solid",
        fgColor=LIGHT_BACKGROUND,
    )
    worksheet["A13"].alignment = Alignment(
        vertical="center",
        wrap_text=True,
    )

    worksheet.row_dimensions[13].height = 24
    worksheet.row_dimensions[14].height = 12

    # Affected worksheets section
    affected_heading_row = 16

    worksheet.merge_cells(
        start_row=affected_heading_row,
        start_column=1,
        end_row=affected_heading_row,
        end_column=6,
    )

    affected_heading_cell = worksheet.cell(
        row=affected_heading_row,
        column=1,
    )

    affected_heading_cell.value = "AFFECTED WORKSHEETS"
    affected_heading_cell.font = Font(
        bold=True,
        color=WHITE,
    )
    affected_heading_cell.fill = PatternFill(
        "solid",
        fgColor=SLATE,
    )

    if audit_result.affected_worksheets:

        for row_number, sheet_name in enumerate(
            audit_result.affected_worksheets,
            start=affected_heading_row + 1,
        ):
            worksheet.merge_cells(
                start_row=row_number,
                start_column=1,
                end_row=row_number,
                end_column=6,
            )

            worksheet.cell(
                row_number,
                1,
                f"• {sheet_name}",
            )

        coverage_heading_row = (
            affected_heading_row
            + len(audit_result.affected_worksheets)
            + 2
        )

    else:
        no_issues_row = affected_heading_row + 1

        worksheet.merge_cells(
            start_row=no_issues_row,
            start_column=1,
            end_row=no_issues_row,
            end_column=6,
        )

        worksheet.cell(
            row=no_issues_row,
            column=1,
            value=(
                "No issues detected by the enabled "
                "audit checks."
            ),
        )

        coverage_heading_row = (
            affected_heading_row + 3
        )

    # Audit coverage section
    worksheet.merge_cells(
        start_row=coverage_heading_row,
        start_column=1,
        end_row=coverage_heading_row,
        end_column=6,
    )

    coverage_heading_cell = worksheet.cell(
        row=coverage_heading_row,
        column=1,
    )

    coverage_heading_cell.value = "AUDIT COVERAGE"
    coverage_heading_cell.font = Font(
        bold=True,
        color=WHITE,
    )
    coverage_heading_cell.fill = PatternFill(
        "solid",
        fgColor=SLATE,
    )

    coverage_status_row = (
        coverage_heading_row + 1
    )

    worksheet.merge_cells(
        start_row=coverage_status_row,
        start_column=1,
        end_row=coverage_status_row,
        end_column=6,
    )

    coverage_status_cell = worksheet.cell(
        row=coverage_status_row,
        column=1,
    )

    coverage_status_cell.value = (
        audit_result.coverage.status
    )

    coverage_status_cell.font = Font(
        size=14,
        bold=True,
        color=NAVY,
    )

    coverage_status_cell.fill = PatternFill(
        "solid",
        fgColor=LIGHT_BACKGROUND,
    )

    coverage_status_cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    worksheet.row_dimensions[
        coverage_status_row
    ].height = 28

    # Layout
    worksheet.column_dimensions["A"].width = 23
    worksheet.column_dimensions["B"].width = 12
    worksheet.column_dimensions["C"].width = 23
    worksheet.column_dimensions["D"].width = 23
    worksheet.column_dimensions["E"].width = 23
    worksheet.column_dimensions["F"].width = 12

    worksheet.row_dimensions[1].height = 25
    worksheet.row_dimensions[2].height = 12
    worksheet.row_dimensions[7].height = 32
    worksheet.row_dimensions[8].height = 22

    worksheet.freeze_panes = None


def _build_findings_sheet(
    worksheet,
    audit_result,
):
    worksheet.sheet_view.showGridLines = False

    headers = [
        "Finding ID",
        "Check ID",
        "Check Name",
        "Worksheet",
        "Location",
        "Priority",
        "Confidence",
    ]

    worksheet.append(headers)

    for finding in _get_priority_sorted_findings(
        audit_result.findings
    ):
        worksheet.append(
            [
                finding.finding_id,
                finding.check_id,
                finding.check_name,
                finding.worksheet,
                finding.location,
                finding.priority,
                finding.confidence,
            ]
        )

    # Header styling
    for cell in worksheet[1]:
        cell.fill = PatternFill(
            "solid",
            fgColor=NAVY,
        )
        cell.font = Font(
            bold=True,
            color=WHITE,
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    worksheet.row_dimensions[1].height = 24

    # Freeze header
    worksheet.freeze_panes = "B2"

    # Add filters
    if worksheet.max_row > 1:
        worksheet.auto_filter.ref = (
            f"A1:G{worksheet.max_row}"
        )

    # Finding row formatting
    for row_number in range(
        2,
        worksheet.max_row + 1,
    ):
        priority = worksheet.cell(
            row=row_number,
            column=6,
        ).value

        # Light alternating row shading
        if row_number % 2 == 0:
            for cell in worksheet[row_number]:
                cell.fill = PatternFill(
                    "solid",
                    fgColor="F8FAFC",
                )

        # Priority highlighting
        priority_cell = worksheet.cell(
            row=row_number,
            column=6,
        )

        if priority == "High Priority":
            priority_cell.fill = PatternFill(
                "solid",
                fgColor=HIGH_PRIORITY_FILL,
            )
            priority_cell.font = Font(
                bold=True,
                color="9C0006",
            )

        elif priority == "Review Recommended":
            priority_cell.fill = PatternFill(
                "solid",
                fgColor=REVIEW_FILL,
            )
            priority_cell.font = Font(
                bold=True,
                color="9C5700",
            )

        elif priority == "Information":
            priority_cell.fill = PatternFill(
                "solid",
                fgColor=INFORMATION_FILL,
            )
            priority_cell.font = Font(
                bold=True,
                color=NAVY,
            )

    # Centre short identifier fields
    for row_number in range(
        2,
        worksheet.max_row + 1,
    ):
        for column_number in (
            1,  # Finding ID
            2,  # Check ID
            7,  # Confidence
        ):
            worksheet.cell(
                row=row_number,
                column=column_number,
            ).alignment = Alignment(
                horizontal="center",
                vertical="top",
            )

    # Subtle table borders
    table_side = Side(
        style="thin",
        color="D9E1E8",
    )

    table_border = Border(
        left=table_side,
        right=table_side,
        top=table_side,
        bottom=table_side,
    )

    for row in worksheet.iter_rows(
        min_row=1,
        max_row=worksheet.max_row,
        min_col=1,
        max_col=7,
    ):
        for cell in row:
            cell.border = table_border
    


def _build_detail_sheet(
    worksheet,
    audit_result,
    ):
    worksheet.sheet_view.showGridLines = False

    headers = [
        "Finding ID",
        "Check ID",
        "Check Name",
        "Priority",
        "Worksheet",
        "Location",
        "Summary",
        "Detail",
        "Why Review",
        "Confidence",
        "Affected Items",
    ]

    worksheet.append(headers)

    for finding in _get_priority_sorted_findings(
        audit_result.findings
    ):
        affected_items_text = _format_affected_items(
            finding.affected_items
        )

        worksheet.append(
            [
                finding.finding_id,
                finding.check_id,
                finding.check_name,
                finding.priority,
                finding.worksheet,
                finding.location,
                finding.summary,
                finding.detail,
                finding.why_review,
                finding.confidence,
                affected_items_text,
            ]
        )


    # Header styling
    for cell in worksheet[1]:
        cell.fill = PatternFill(
            "solid",
            fgColor=NAVY,
        )
        cell.font = Font(
            bold=True,
            color=WHITE,
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    worksheet.row_dimensions[1].height = 28

    # Freeze header and enable filters
    worksheet.freeze_panes = "B2"

    if worksheet.max_row > 1:
        worksheet.auto_filter.ref = (
            f"A1:K{worksheet.max_row}"
        )

    # Subtle borders
    table_side = Side(
        style="thin",
        color="D9E1E8",
    )

    table_border = Border(
        left=table_side,
        right=table_side,
        top=table_side,
        bottom=table_side,
    )

    for row in worksheet.iter_rows(
        min_row=1,
        max_row=worksheet.max_row,
        min_col=1,
        max_col=11,
    ):
        for cell in row:
            cell.border = table_border
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    # Finding formatting
    for row_number in range(
        2,
        worksheet.max_row + 1,
    ):
        # Alternating row background
        if row_number % 2 == 0:
            for cell in worksheet[row_number]:
                cell.fill = PatternFill(
                    "solid",
                    fgColor="F8FAFC",
                )

        priority_cell = worksheet.cell(
            row=row_number,
            column=4,
        )

        priority = priority_cell.value

        if priority == "High Priority":
            priority_cell.fill = PatternFill(
                "solid",
                fgColor=HIGH_PRIORITY_FILL,
            )
            priority_cell.font = Font(
                bold=True,
                color="9C0006",
            )

        elif priority == "Review Recommended":
            priority_cell.fill = PatternFill(
                "solid",
                fgColor=REVIEW_FILL,
            )
            priority_cell.font = Font(
                bold=True,
                color="9C5700",
            )

        elif priority == "Information":
            priority_cell.fill = PatternFill(
                "solid",
                fgColor=INFORMATION_FILL,
            )
            priority_cell.font = Font(
                bold=True,
                color=NAVY,
            )

        # Emphasise finding summary
        worksheet.cell(
            row=row_number,
            column=7,
        ).font = Font(
            bold=True,
            color=NAVY,
        )

        # Centre short fields
        for column_number in (
            1,   # Finding ID
            2,   # Check ID
            6,   # Location
            10,  # Confidence
        ):
            worksheet.cell(
                row=row_number,
                column=column_number,
            ).alignment = Alignment(
                horizontal="center",
                vertical="top",
                wrap_text=True,
            )


def _format_affected_items(affected_items):
    """
    Convert structured finding evidence into readable
    report text without changing the underlying model.
    """

    if not affected_items:
        return ""

    formatted_items = []

    for item in affected_items:
        parts = []

        for key, value in item.items():
            if value is None:
                continue

            readable_key = key.replace(
                "_",
                " ",
            ).title()

            parts.append(
                f"{readable_key}: {value}"
            )

        formatted_items.append(
            " | ".join(parts)
        )

    return "\n".join(formatted_items)


def _build_information_sheet(
    worksheet,
    audit_result,
):
    worksheet.sheet_view.showGridLines = False

    # Title
    worksheet.merge_cells("A1:D2")
    worksheet["A1"] = "AUDIT INFORMATION"
    worksheet["A1"].font = Font(
        size=18,
        bold=True,
        color=WHITE,
    )
    worksheet["A1"].fill = PatternFill(
        "solid",
        fgColor=NAVY,
    )
    worksheet["A1"].alignment = Alignment(
        vertical="center",
    )

    # Audit details
    information = [
        (
            "Workbook",
            audit_result.workbook_name,
        ),
        (
            "Audit Version",
            "V1",
        ),
        (
            "Checks Enabled",
            audit_result.checks_completed,
        ),
        (
            "Total Findings",
            audit_result.total_findings,
        ),
        (
            "Scope",
            (
                "Formula, workbook structure and "
                "tabular data quality checks"
            ),
        ),
    ]

    for row_number, (label, value) in enumerate(
        information,
        start=4,
    ):
        worksheet.cell(
            row=row_number,
            column=1,
            value=label,
        )

        worksheet.cell(
            row=row_number,
            column=2,
            value=value,
        )

        worksheet.cell(
            row=row_number,
            column=1,
        ).font = Font(
            bold=True,
            color=SLATE,
        )

        worksheet.cell(
            row=row_number,
            column=1,
        ).fill = PatternFill(
            "solid",
            fgColor=LIGHT_BACKGROUND,
        )

        worksheet.cell(
            row=row_number,
            column=2,
        ).alignment = Alignment(
            horizontal="left",
            vertical="top",
            wrap_text=True,
        )

    # Audit coverage section
    worksheet.merge_cells("A10:D10")
    worksheet["A10"] = "AUDIT COVERAGE"
    worksheet["A10"].font = Font(
        bold=True,
        color=WHITE,
    )
    worksheet["A10"].fill = PatternFill(
        "solid",
        fgColor=SLATE,
    )

    check_names = {
        "EA01": "Broken Formula Reference",
        "EA02": "Formula Pattern Inconsistency",
        "EA03": "Hardcoded Formula Override",
        "EA04": "Formula Error Detection",
        "EA05": "Hidden Calculation Sheet",
        "EA06": "Suspicious Blank Detection",
        "EA07": "Inconsistent Data Type",
        "EA08": "Duplicate Record Detection",
    }

    affected_checks = (
        ", ".join(
            (
                f"{check_names.get(check_id, check_id)} "
                f"({check_id})"
            )
            for check_id
            in audit_result.coverage.affected_checks
        )
        if audit_result.coverage.affected_checks
        else "None"
    )

    coverage_information = [
        (
            "Status",
            audit_result.coverage.status,
        ),
        (
            "Formula Cells Assessed",
            audit_result.coverage.formula_cells,
        ),
        (
            "Missing Cached Results",
            audit_result.coverage.missing_cached_results,
        ),
        (
            "Checks with Reduced Coverage",
            affected_checks,
        ),
    ]

    for row_number, (label, value) in enumerate(
        coverage_information,
        start=11,
    ):
        worksheet.cell(
            row=row_number,
            column=1,
            value=label,
        )

        worksheet.cell(
            row=row_number,
            column=2,
            value=value,
        )

        worksheet.cell(
            row=row_number,
            column=1,
        ).font = Font(
            bold=True,
            color=SLATE,
        )

        worksheet.cell(
            row=row_number,
            column=1,
        ).fill = PatternFill(
            "solid",
            fgColor=LIGHT_BACKGROUND,
        )

        worksheet.cell(
            row=row_number,
            column=2,
        ).alignment = Alignment(
            horizontal="left",
            vertical="top",
            wrap_text=True,
        )

    # Emphasise coverage status
    worksheet["B11"].font = Font(
        bold=True,
        color=NAVY,
    )

    # Coverage warning
    if audit_result.coverage.warnings:
        worksheet.merge_cells("A16:D16")
        worksheet["A16"] = "COVERAGE WARNING"
        worksheet["A16"].font = Font(
            bold=True,
            color="9C5700",
        )
        worksheet["A16"].fill = PatternFill(
            "solid",
            fgColor=REVIEW_FILL,
        )

        worksheet.merge_cells("A17:D18")
        worksheet["A17"] = "\n".join(
            audit_result.coverage.warnings
        )
        worksheet["A17"].alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )
        worksheet["A17"].fill = PatternFill(
            "solid",
            fgColor="FFF8E7",
        )

        important_heading_row = 20
        important_text_row = 21

    else:
        important_heading_row = 16
        important_text_row = 17

    # Important information section
    worksheet.merge_cells(
        start_row=important_heading_row,
        start_column=1,
        end_row=important_heading_row,
        end_column=4,
    )

    important_heading_cell = worksheet.cell(
        row=important_heading_row,
        column=1,
    )

    important_heading_cell.value = "IMPORTANT"
    important_heading_cell.font = Font(
        bold=True,
        color=WHITE,
    )
    important_heading_cell.fill = PatternFill(
        "solid",
        fgColor=SLATE,
    )

    worksheet.merge_cells(
        start_row=important_text_row,
        start_column=1,
        end_row=important_text_row + 1,
        end_column=4,
    )

    important_text_cell = worksheet.cell(
        row=important_text_row,
        column=1,
    )

    important_text_cell.value = (
        "This report identifies issues detected by the "
        "enabled audit checks. It does not guarantee that "
        "the source workbook is free from all errors."
    )

    important_text_cell.alignment = Alignment(
        vertical="top",
        wrap_text=True,
    )

    important_text_cell.fill = PatternFill(
        "solid",
        fgColor=LIGHT_BACKGROUND,
    )

    # Layout
    worksheet.column_dimensions["A"].width = 30
    worksheet.column_dimensions["B"].width = 40
    worksheet.column_dimensions["C"].width = 18
    worksheet.column_dimensions["D"].width = 18

    worksheet.row_dimensions[1].height = 24
    worksheet.row_dimensions[
        important_text_row
    ].height = 28

    worksheet.freeze_panes = None


def _apply_basic_formatting(workbook):
    """
    Apply consistent V1 formatting across the report.

    Retained temporarily while individual report sheets
    receive their own presentation rules.
    """

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for worksheet in workbook.worksheets:

        worksheet.freeze_panes = "A2"

        for cell in worksheet[1]:
            if cell.value is not None:
                cell.fill = header_fill
                cell.font = header_font

        for column_cells in worksheet.columns:
            max_length = 0

            column_letter = get_column_letter(
                column_cells[0].column
            )

            for cell in column_cells:
                if cell.value is None:
                    continue

                value_length = len(
                    str(cell.value)
                )

                max_length = max(
                    max_length,
                    value_length,
                )

                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

            worksheet.column_dimensions[
                column_letter
            ].width = min(
                max(max_length + 2, 12),
                50,
            )


def _adjust_detail_row_heights(worksheet):
    """
    Estimate row heights for wrapped Finding Detail text.
    """

    for row_number in range(
        2,
        worksheet.max_row + 1,
    ):
        narrative_columns = (
            7,   # Summary
            8,   # Detail
            9,   # Why Review
            11,  # Affected Items
        )

        longest_text = max(
            (
                len(
                    str(
                        worksheet.cell(
                            row=row_number,
                            column=column_number,
                        ).value
                        or ""
                    )
                )
                for column_number
                in narrative_columns
            ),
            default=0,
        )

        estimated_lines = max(
            1,
            (longest_text // 40) + 1,
        )

        worksheet.row_dimensions[
            row_number
        ].height = min(
            max(30, estimated_lines * 15),
            90,
        )


def _autofit_and_wrap(
    worksheet,
    min_width=10,
    max_width=45,
):
    """
    Apply sensible automatic column widths and wrap
    longer text without creating excessively wide columns.
    """

    for column_cells in worksheet.columns:
        column_letter = get_column_letter(
            column_cells[0].column
        )

        max_length = 0

        for cell in column_cells:
            if cell.value is None:
                continue

            text = str(cell.value)

            # Use the longest line when a cell
            # contains line breaks.
            longest_line = max(
                len(line)
                for line in text.split("\n")
            )

            max_length = max(
                max_length,
                longest_line,
            )

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

        calculated_width = max(
            min_width,
            min(max_length + 2, max_width),
        )

        worksheet.column_dimensions[
            column_letter
        ].width = calculated_width


def save_audit_report(
    audit_result,
    output_path,
):
    """
    Build and save an Excel audit report.
    """

    workbook = build_audit_report(
        audit_result
    )

    workbook.save(output_path)

    return Path(output_path)