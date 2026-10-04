from audit_checks import check_formula_pattern_inconsistencies
from workbook_utils import load_audit_workbook


def test_ea02_detects_planted_formula_inconsistency():
    workbook = load_audit_workbook(
        "data/Apex_FY26_Forecast.xlsx"
    )

    findings = check_formula_pattern_inconsistencies(workbook)

    matches = [
        finding
        for finding in findings
        if finding.worksheet == "Sales Forecast"
        and finding.location == "E9"
    ]

    assert len(matches) == 1
    assert matches[0].check_id == "EA02"
    assert matches[0].priority == "Review Recommended"
    assert matches[0].confidence == "Medium"


def test_ea02_returns_only_expected_apex_finding():
    workbook = load_audit_workbook(
        "data/Apex_FY26_Forecast.xlsx"
    )

    findings = check_formula_pattern_inconsistencies(workbook)

    locations = {
        (finding.worksheet, finding.location)
        for finding in findings
    }

    assert locations == {
        ("Sales Forecast", "E9")
    }


def test_ea02_clean_workbook_has_no_formula_pattern_findings():
    workbook = load_audit_workbook(
        "data/Apex_FY26_Forecast_CLEAN.xlsx"
    )

    findings = check_formula_pattern_inconsistencies(workbook)

    assert len(findings) == 0


def test_ea02_does_not_compare_unrelated_formula_families(
    tmp_path,
):
    from openpyxl import Workbook

    from audit_checks import (
        check_formula_pattern_inconsistencies,
    )

    file_path = (
        tmp_path
        / "mixed_formula_families.xlsx"
    )

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Analysis"

    # Date calculation family
    worksheet["B2"] = "=EOMONTH(B1,0)"
    worksheet["C2"] = "=EOMONTH(C1,0)"
    worksheet["D2"] = "=EOMONTH(D1,0)"
    worksheet["E2"] = "=EOMONTH(E1,0)"

    # Revenue calculation family
    worksheet["B3"] = "=B10*B11"
    worksheet["C3"] = "=C10*C11"
    worksheet["D3"] = "=D10*D11"
    worksheet["E3"] = "=E10*E11"

    # Expense calculation family
    worksheet["B4"] = "=B12+B13"
    worksheet["C4"] = "=C12+C13"
    worksheet["D4"] = "=D12+D13"
    worksheet["E4"] = "=E12+E13"

    workbook.save(
        file_path
    )

    audit_workbook = load_audit_workbook(
        file_path
    )

    findings = check_formula_pattern_inconsistencies(
        audit_workbook
    )

    assert findings == []


def test_ea02_detects_outlier_inside_repeated_formula_family(
    tmp_path,
):
    from openpyxl import Workbook

    file_path = (
        tmp_path
        / "formula_outlier.xlsx"
    )

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Forecast"

    worksheet["B2"] = "=C2*D2"
    worksheet["B3"] = "=C3*D3"
    worksheet["B4"] = "=C4*D4"

    # Intentional anomaly
    worksheet["B5"] = "=C5+D5"

    worksheet["B6"] = "=C6*D6"
    worksheet["B7"] = "=C7*D7"

    workbook.save(
        file_path
    )

    audit_workbook = load_audit_workbook(
        file_path
    )

    findings = check_formula_pattern_inconsistencies(
        audit_workbook
    )

    locations = [
        finding.location
        for finding in findings
    ]

    assert locations == ["B5"]


def test_ea02_does_not_treat_formula_family_boundary_as_anomaly(
    tmp_path,
):
    from openpyxl import Workbook

    from audit_checks import (
        check_formula_pattern_inconsistencies,
    )

    file_path = (
        tmp_path
        / "formula_family_boundary.xlsx"
    )

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Analysis"

    # First legitimate family
    worksheet["B2"] = "=B10*B11"
    worksheet["C2"] = "=C10*C11"
    worksheet["D2"] = "=D10*D11"
    worksheet["E2"] = "=E10*E11"

    # Second legitimate family
    worksheet["B3"] = "=SUM(B20:B25)"
    worksheet["C3"] = "=SUM(C20:C25)"
    worksheet["D3"] = "=SUM(D20:D25)"
    worksheet["E3"] = "=SUM(E20:E25)"

    workbook.save(
        file_path
    )

    audit_workbook = load_audit_workbook(
        file_path
    )

    findings = check_formula_pattern_inconsistencies(
        audit_workbook
    )

    assert findings == []


def test_ea02_does_not_flag_multiple_legitimate_vertical_formula_families(
    tmp_path,
):
    from openpyxl import Workbook

    file_path = (
        tmp_path
        / "multiple_vertical_formula_families.xlsx"
    )

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Analysis"

    # Formula family 1
    worksheet["B2"] = "=C2*D2"
    worksheet["B3"] = "=C3*D3"
    worksheet["B4"] = "=C4*D4"
    worksheet["B5"] = "=C5*D5"
    worksheet["B6"] = "=C6*D6"

    # Formula family 2
    worksheet["B7"] = "=SUM(C7:D7)"
    worksheet["B8"] = "=SUM(C8:D8)"
    worksheet["B9"] = "=SUM(C9:D9)"
    worksheet["B10"] = "=SUM(C10:D10)"

    # Formula family 3
    worksheet["B11"] = "=IF(C11>0,C11,0)"
    worksheet["B12"] = "=IF(C12>0,C12,0)"
    worksheet["B13"] = "=IF(C13>0,C13,0)"
    worksheet["B14"] = "=IF(C14>0,C14,0)"

    workbook.save(
        file_path
    )

    audit_workbook = load_audit_workbook(
        file_path
    )

    findings = check_formula_pattern_inconsistencies(
        audit_workbook
    )

    assert findings == []