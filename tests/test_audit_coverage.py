from openpyxl import Workbook

from audit_coverage import assess_audit_coverage


def test_complete_coverage_when_cached_results_available():
    result = assess_audit_coverage(
        "data/Apex_FY26_Forecast_CLEAN.xlsx"
    )

    assert result.status == "Complete"
    assert result.formula_cells == 43
    assert result.missing_cached_results == 0
    assert result.affected_checks == []
    assert result.warnings == []


def test_limited_coverage_when_cached_results_missing(
    tmp_path,
):
    test_file = tmp_path / "missing_cache.xlsx"

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Test Data"

    worksheet["A1"] = 100
    worksheet["B1"] = 200
    worksheet["C1"] = "=A1+B1"

    workbook.save(test_file)

    result = assess_audit_coverage(
        test_file
    )

    assert result.status == "Limited"
    assert result.formula_cells == 1
    assert result.missing_cached_results == 1

    assert result.affected_checks == [
        "EA04",
    ]

    assert len(result.warnings) == 1
    assert "1 formula cell(s)" in result.warnings[0]


def test_ea08_limited_when_missing_cache_is_in_data_region(
    tmp_path,
):
    test_file = tmp_path / "missing_table_cache.xlsx"

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Orders"

    worksheet.append(
        [
            "Order ID",
            "Quantity",
            "Unit Price",
            "Total",
        ]
    )

    for row_number in range(
        2,
        12,
    ):
        worksheet.cell(
            row=row_number,
            column=1,
            value=f"ORD{row_number:03d}",
        )

        worksheet.cell(
            row=row_number,
            column=2,
            value=10,
        )

        worksheet.cell(
            row=row_number,
            column=3,
            value=5,
        )

        worksheet.cell(
            row=row_number,
            column=4,
            value=(
                f"=B{row_number}*C{row_number}"
            ),
        )

    workbook.save(test_file)

    result = assess_audit_coverage(
        test_file
    )

    assert result.status == "Limited"
    assert result.formula_cells == 10
    assert result.missing_cached_results == 10

    assert result.affected_checks == [
        "EA04",
        "EA08",
    ]