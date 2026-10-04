from audit_checks import check_inconsistent_data_types
from workbook_utils import load_audit_workbook


def test_ea07_detects_planted_type_anomaly():
    workbook = load_audit_workbook(
        "data/Apex_FY26_Forecast.xlsx"
    )

    findings = check_inconsistent_data_types(workbook)

    matches = [
        finding
        for finding in findings
        if finding.worksheet == "Customer Orders"
        and finding.location == "B15"
    ]

    assert len(matches) == 1
    assert matches[0].check_id == "EA07"
    assert matches[0].priority == "Review Recommended"
    assert matches[0].confidence == "Medium"


def test_ea07_returns_only_expected_apex_finding():
    workbook = load_audit_workbook(
        "data/Apex_FY26_Forecast.xlsx"
    )

    findings = check_inconsistent_data_types(workbook)

    locations = {
        (finding.worksheet, finding.location)
        for finding in findings
    }

    assert locations == {
        ("Customer Orders", "B15")
    }


def test_ea07_does_not_misclassify_formula_override():
    workbook = load_audit_workbook(
        "data/Apex_FY26_Forecast.xlsx"
    )

    findings = check_inconsistent_data_types(workbook)

    locations = {
        (finding.worksheet, finding.location)
        for finding in findings
    }

    assert ("Sales Forecast", "E14") not in locations


def test_ea07_does_not_flag_product_identifiers():
    workbook = load_audit_workbook(
        "data/Apex_FY26_Forecast.xlsx"
    )

    findings = check_inconsistent_data_types(workbook)

    product_master_findings = [
        finding
        for finding in findings
        if finding.worksheet == "Product Master"
    ]

    assert len(product_master_findings) == 0


def test_ea07_does_not_treat_formula_summary_as_header(tmp_path):
    from openpyxl import Workbook

    file_path = tmp_path / "formula_summary.xlsx"

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Budget"

    # Summary labels.
    worksheet["A1"] = "Current"
    worksheet["B1"] = "Projected"
    worksheet["C1"] = "Actual"

    # Summary formulas must not be treated as table headers.
    worksheet["A2"] = "=SUM(A5:A20)"
    worksheet["B2"] = "=SUM(B5:B20)"
    worksheet["C2"] = "=SUM(C5:C20)"

    # Blank separator.
    # Row 3 intentionally left blank.

    # Genuine table header.
    worksheet["A4"] = "Category"
    worksheet["B4"] = "Projected Average"
    worksheet["C4"] = "Actual Average"

    # Genuine table data.
    for row in range(5, 21):
        worksheet.cell(row, 1).value = f"Expense {row}"
        worksheet.cell(row, 2).value = row * 5
        worksheet.cell(row, 3).value = row * 4

    workbook.save(file_path)

    audit_workbook = load_audit_workbook(file_path)

    findings = check_inconsistent_data_types(
        audit_workbook
    )

    affected_locations = {
        item["location"]
        for finding in findings
        for item in finding.affected_items
    }

    assert "B4" not in affected_locations
    assert "C4" not in affected_locations