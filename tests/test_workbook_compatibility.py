from pathlib import Path
import re
import shutil
import zipfile
import hashlib

from openpyxl import Workbook, load_workbook

from workbook_compatibility import (
    create_compatible_audit_copy,
)


def test_compatible_workbook_requires_no_adjustment():
    result = create_compatible_audit_copy(
        "data/Apex_FY26_Forecast.xlsx"
    )

    try:
        assert result["adjustment_applied"] is False
        assert result["adjustments"] == []

        compatible_path = Path(
            result["path"]
        )

        assert compatible_path.exists()
        assert (
            compatible_path.name
            == "Apex_FY26_Forecast.xlsx"
        )

    finally:
        shutil.rmtree(
            result["temporary_directory"],
            ignore_errors=True,
        )


def test_unsupported_font_family_is_normalised(
    tmp_path,
):
    source_path = (
        tmp_path
        / "unsupported_font_family.xlsx"
    )

    normal_path = (
        tmp_path
        / "normal.xlsx"
    )

    # Create a normal workbook first.
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["A1"] = "Test"
    workbook.save(normal_path)
    workbook.close()

    # Rebuild the XLSX package with an unsupported
    # font-family value in styles.xml.
    with zipfile.ZipFile(
        normal_path,
        "r",
    ) as source_archive:
        with zipfile.ZipFile(
            source_path,
            "w",
            zipfile.ZIP_DEFLATED,
        ) as target_archive:

            for item in source_archive.infolist():
                content = source_archive.read(
                    item.filename
                )

                if item.filename == "xl/styles.xml":
                    content, replacements = re.subn(
                        rb'(<family[^>]*\bval=")2("[^>]*/?>)',
                        rb'\g<1>34\g<2>',
                        content,
                        count=1,
                    )

                    # Make sure the test fixture was
                    # actually made incompatible.
                    assert replacements == 1

                target_archive.writestr(
                    item,
                    content,
                )

    result = create_compatible_audit_copy(
        source_path
    )

    try:
        assert (
            result["adjustment_applied"]
            is True
        )

        assert result["adjustments"]

        compatible_path = Path(
            result["path"]
        )

        assert compatible_path.exists()

        # The repaired workbook must be readable.
        repaired_workbook = load_workbook(
            compatible_path
        )

        assert (
            repaired_workbook.active["A1"].value
            == "Test"
        )

        repaired_workbook.close()

        # Verify that the unsupported font-family
        # metadata has been removed.
        with zipfile.ZipFile(
            compatible_path,
            "r",
        ) as archive:
            styles_xml = archive.read(
                "xl/styles.xml"
            )

        assert re.search(
            rb'<family[^>]*\bval="34"[^>]*/?>',
            styles_xml,
        ) is None

        assert re.search(
            rb'<family[^>]*\bval="2"[^>]*/?>',
            styles_xml,
        ) is not None

    finally:
        shutil.rmtree(
            result["temporary_directory"],
            ignore_errors=True,
        )


def test_compatibility_process_does_not_modify_source(
    tmp_path,
):
    source_path = (
        tmp_path
        / "source_workbook.xlsx"
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet["A1"] = "Original Content"
    worksheet["B1"] = 12345
    worksheet["C1"] = "=B1*2"
    workbook.save(source_path)
    workbook.close()

    before_hash = hashlib.sha256(
        source_path.read_bytes()
    ).hexdigest()

    result = create_compatible_audit_copy(
        source_path
    )

    try:
        after_hash = hashlib.sha256(
            source_path.read_bytes()
        ).hexdigest()

        assert before_hash == after_hash

        assert (
            Path(result["path"])
            != source_path
        )

        assert source_path.exists()

    finally:
        shutil.rmtree(
            result["temporary_directory"],
            ignore_errors=True,
        )


def test_repair_preserves_workbook_content(
    tmp_path,
):
    normal_path = (
        tmp_path
        / "normal_content.xlsx"
    )

    source_path = (
        tmp_path
        / "incompatible_content.xlsx"
    )

    # Create a workbook containing representative
    # business content.
    workbook = Workbook()

    summary = workbook.active
    summary.title = "Summary"

    summary["A1"] = "Revenue"
    summary["B1"] = 125000
    summary["C1"] = "=B1*1.2"

    detail = workbook.create_sheet(
        "Detail"
    )

    detail["A1"] = "Product"
    detail["B1"] = "Quantity"
    detail["C1"] = "Unit Price"
    detail["D1"] = "Total"

    detail["A2"] = "Part A"
    detail["B2"] = 10
    detail["C2"] = 25
    detail["D2"] = "=B2*C2"

    detail["A3"] = "Part B"
    detail["B3"] = 5
    detail["C3"] = 40
    detail["D3"] = "=B3*C3"

    workbook.save(normal_path)
    workbook.close()

    # Create an incompatible copy by injecting an
    # unsupported font-family value.
    with zipfile.ZipFile(
        normal_path,
        "r",
    ) as source_archive:
        with zipfile.ZipFile(
            source_path,
            "w",
            zipfile.ZIP_DEFLATED,
        ) as target_archive:

            for item in source_archive.infolist():
                content = source_archive.read(
                    item.filename
                )

                if item.filename == "xl/styles.xml":
                    content, replacements = re.subn(
                        rb'(<family[^>]*\bval=")2("[^>]*/?>)',
                        rb'\g<1>34\g<2>',
                        content,
                        count=1,
                    )

                    assert replacements == 1

                target_archive.writestr(
                    item,
                    content,
                )

    result = create_compatible_audit_copy(
        source_path
    )

    try:
        assert (
            result["adjustment_applied"]
            is True
        )

        repaired_workbook = load_workbook(
            result["path"],
            data_only=False,
        )

        # Worksheet structure is preserved.
        assert repaired_workbook.sheetnames == [
            "Summary",
            "Detail",
        ]

        # Values are preserved.
        assert (
            repaired_workbook["Summary"]["A1"].value
            == "Revenue"
        )
        assert (
            repaired_workbook["Summary"]["B1"].value
            == 125000
        )

        assert (
            repaired_workbook["Detail"]["A2"].value
            == "Part A"
        )
        assert (
            repaired_workbook["Detail"]["B3"].value
            == 5
        )
        assert (
            repaired_workbook["Detail"]["C3"].value
            == 40
        )

        # Formulas are preserved exactly.
        assert (
            repaired_workbook["Summary"]["C1"].value
            == "=B1*1.2"
        )
        assert (
            repaired_workbook["Detail"]["D2"].value
            == "=B2*C2"
        )
        assert (
            repaired_workbook["Detail"]["D3"].value
            == "=B3*C3"
        )

        repaired_workbook.close()

    finally:
        shutil.rmtree(
            result["temporary_directory"],
            ignore_errors=True,
        )


def test_repair_changes_only_affected_xml_component(
    tmp_path,
):
    normal_path = (
        tmp_path
        / "normal_package.xlsx"
    )

    source_path = (
        tmp_path
        / "incompatible_package.xlsx"
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Data"

    worksheet["A1"] = "Quantity"
    worksheet["B1"] = "Price"
    worksheet["C1"] = "Total"

    worksheet["A2"] = 10
    worksheet["B2"] = 25
    worksheet["C2"] = "=A2*B2"

    workbook.save(normal_path)
    workbook.close()

    # Create an incompatible source workbook by
    # changing only xl/styles.xml.
    with zipfile.ZipFile(
        normal_path,
        "r",
    ) as source_archive:
        with zipfile.ZipFile(
            source_path,
            "w",
            zipfile.ZIP_DEFLATED,
        ) as target_archive:

            for item in source_archive.infolist():
                content = source_archive.read(
                    item.filename
                )

                if item.filename == "xl/styles.xml":
                    content, replacements = re.subn(
                        rb'(<family[^>]*\bval=")2("[^>]*/?>)',
                        rb'\g<1>34\g<2>',
                        content,
                        count=1,
                    )

                    assert replacements == 1

                target_archive.writestr(
                    item,
                    content,
                )

    result = create_compatible_audit_copy(
        source_path
    )

    try:
        assert (
            result["adjustment_applied"]
            is True
        )

        repaired_path = Path(
            result["path"]
        )

        with zipfile.ZipFile(
            source_path,
            "r",
        ) as source_archive:
            source_files = {
                name: source_archive.read(name)
                for name in source_archive.namelist()
            }

        with zipfile.ZipFile(
            repaired_path,
            "r",
        ) as repaired_archive:
            repaired_files = {
                name: repaired_archive.read(name)
                for name in repaired_archive.namelist()
            }

        # No package components should be added
        # or removed during compatibility repair.
        assert (
            set(source_files)
            == set(repaired_files)
        )

        changed_components = {
            name
            for name in source_files
            if (
                source_files[name]
                != repaired_files[name]
            )
        }

        # Only the XML component containing the
        # unsupported metadata should change.
        assert changed_components == {
            "xl/styles.xml"
        }

    finally:
        shutil.rmtree(
            result["temporary_directory"],
            ignore_errors=True,
        )