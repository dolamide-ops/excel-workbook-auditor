from audit_checks import check_formula_errors
from workbook_utils import (
    load_audit_workbook,
    load_cached_workbook,
)


def test_ea04_detects_divide_by_zero_error():
    path = "data/Apex_FY26_Forecast.xlsx"

    formula_workbook = load_audit_workbook(path)
    cached_workbook = load_cached_workbook(path)

    findings = check_formula_errors(
        formula_workbook,
        cached_workbook,
    )

    matches = [
        finding
        for finding in findings
        if finding.worksheet == "Sales Forecast"
        and finding.location == "J6"
    ]

    assert len(matches) == 1
    assert matches[0].check_id == "EA04"
    assert matches[0].priority == "High Priority"
    assert matches[0].confidence == "High"


def test_ea04_detects_na_as_review_recommended():
    path = "data/Apex_FY26_Forecast.xlsx"

    formula_workbook = load_audit_workbook(path)
    cached_workbook = load_cached_workbook(path)

    findings = check_formula_errors(
        formula_workbook,
        cached_workbook,
    )

    matches = [
        finding
        for finding in findings
        if finding.worksheet == "Product Master"
        and finding.location == "I7"
    ]

    assert len(matches) == 1
    assert matches[0].priority == "Review Recommended"
    assert matches[0].confidence == "High"


def test_ea04_does_not_duplicate_ea01_ref_error():
    path = "data/Apex_FY26_Forecast.xlsx"

    formula_workbook = load_audit_workbook(path)
    cached_workbook = load_cached_workbook(path)

    findings = check_formula_errors(
        formula_workbook,
        cached_workbook,
    )

    locations = {
        (finding.worksheet, finding.location)
        for finding in findings
    }

    assert ("Executive Summary", "B10") not in locations


def test_ea04_returns_only_expected_apex_findings():
    path = "data/Apex_FY26_Forecast.xlsx"

    formula_workbook = load_audit_workbook(path)
    cached_workbook = load_cached_workbook(path)

    findings = check_formula_errors(
        formula_workbook,
        cached_workbook,
    )

    locations = {
        (finding.worksheet, finding.location)
        for finding in findings
    }

    assert locations == {
        ("Sales Forecast", "J6"),
        ("Product Master", "I7"),
    }


def test_ea04_groups_contiguous_same_error_cells():
    from openpyxl import Workbook

    formula_workbook = Workbook()
    formula_sheet = formula_workbook.active
    formula_sheet.title = "Analysis"

    cached_workbook = Workbook()
    cached_sheet = cached_workbook.active
    cached_sheet.title = "Analysis"

    for row in range(2, 5):
        for column in range(2, 6):
            formula_cell = formula_sheet.cell(
                row=row,
                column=column,
            )
            formula_cell.value = "=NA()"

            cached_cell = cached_sheet.cell(
                row=row,
                column=column,
            )
            cached_cell.value = "#N/A"

    findings = check_formula_errors(
        formula_workbook,
        cached_workbook,
    )

    assert len(findings) == 1

    finding = findings[0]

    assert finding.check_id == "EA04"
    assert finding.priority == "Review Recommended"
    assert finding.confidence == "High"

    assert len(finding.affected_items) == 12

    affected_locations = {
        item["location"]
        for item in finding.affected_items
    }

    assert affected_locations == {
        "B2", "C2", "D2", "E2",
        "B3", "C3", "D3", "E3",
        "B4", "C4", "D4", "E4",
    }


def test_ea04_keeps_separate_error_blocks_as_separate_findings():
    from openpyxl import Workbook

    formula_workbook = Workbook()
    formula_sheet = formula_workbook.active
    formula_sheet.title = "Analysis"

    cached_workbook = Workbook()
    cached_sheet = cached_workbook.active
    cached_sheet.title = "Analysis"

    # First #N/A block
    for row in range(2, 4):
        for column in range(2, 4):
            formula_sheet.cell(
                row=row,
                column=column,
            ).value = "=NA()"

            cached_sheet.cell(
                row=row,
                column=column,
            ).value = "#N/A"

    # Second #N/A block, separated by a blank row
    for row in range(5, 7):
        for column in range(2, 4):
            formula_sheet.cell(
                row=row,
                column=column,
            ).value = "=NA()"

            cached_sheet.cell(
                row=row,
                column=column,
            ).value = "#N/A"

    findings = check_formula_errors(
        formula_workbook,
        cached_workbook,
    )

    assert len(findings) == 2

    affected_counts = sorted(
        len(finding.affected_items)
        for finding in findings
    )

    assert affected_counts == [4, 4]