from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator

from datetime import date, datetime


EXCEL_ERROR_VALUES = {
    "#NULL!",
    "#DIV/0!",
    "#VALUE!",
    "#REF!",
    "#NAME?",
    "#NUM!",
    "#N/A",
}


IDENTIFIER_TERMS = {
    "id",
    "code",
    "number",
    "no",
    "reference",
    "ref",
    "key",
}


def is_identifier_header(header: str) -> bool:
    """
    Identify columns whose values primarily represent identifiers
    rather than quantities or measures.
    """
    if not isinstance(header, str):
        return False

    normalised = (
        header.lower()
        .replace("_", " ")
        .replace("-", " ")
        .strip()
    )

    words = set(normalised.split())

    return bool(words & IDENTIFIER_TERMS)


def load_audit_workbook(file_path: str | Path):
    """
    Load an Excel workbook while preserving formulas.

    The auditor must inspect formula text, so data_only=False
    is required for the primary workbook representation.
    """
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"Workbook not found: {path}")

    if path.suffix.lower() != ".xlsx":
        raise ValueError("Only .xlsx workbooks are supported in V1.")

    return load_workbook(path, data_only=False)


def load_cached_workbook(file_path: str | Path):
    """
    Load the workbook using Excel's last saved cached formula results.

    openpyxl does not calculate formulas itself, so cached values
    may be unavailable depending on how the workbook was last saved.
    """
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"Workbook not found: {path}")

    if path.suffix.lower() != ".xlsx":
        raise ValueError("Only .xlsx workbooks are supported in V1.")

    return load_workbook(path, data_only=True)


from openpyxl.formula.translate import Translator



def normalise_formula(cell):
    """
    Translate a formula to a safe common reference position so formulas
    with the same relative structure can be compared.

    A distant anchor is used to prevent relative references from being
    translated outside Excel's valid row/column range.

    Example:
        E6  =C6*D6
        E7  =C7*D7

    Both normalise to the same relative formula pattern.
    """
    if cell.data_type != "f" or not isinstance(cell.value, str):
        return None

    try:
        return Translator(
            cell.value,
            origin=cell.coordinate
        ).translate_formula("ZZ1000")

    except Exception:
        return None


def formula_references_sheet(formula: str, sheet_name: str) -> bool:
    """
    Return True when an Excel formula contains a reference
    to the supplied worksheet.

    Handles common references such as:
        =Calc_Helper!B7
        ='Calc Helper'!B7
    """
    if not isinstance(formula, str):
        return False

    quoted_reference = f"'{sheet_name}'!"
    unquoted_reference = f"{sheet_name}!"

    formula_upper = formula.upper()

    return (
        quoted_reference.upper() in formula_upper
        or unquoted_reference.upper() in formula_upper
    )


def detect_data_regions(worksheet):
    """
    Detect simple table-like data regions in a worksheet.

    V1 looks for rows containing at least three populated cells
    followed by at least three populated data rows.
    """
    from audit_models import DataRegion

    regions = []

    max_row = worksheet.max_row
    max_column = worksheet.max_column

    for row_number in range(1, max_row + 1):

        row_values = [
            worksheet.cell(row_number, column).value
            for column in range(1, max_column + 1)
        ]

        populated_columns = [
            index + 1
            for index, value in enumerate(row_values)
            if value not in (None, "")
        ]

        if len(populated_columns) < 3:
            continue

        start_column = min(populated_columns)
        end_column = max(populated_columns)

        headers = [
            worksheet.cell(row_number, column).value
            for column in range(start_column, end_column + 1)
        ]

        # Header candidates should predominantly contain text.
        text_headers = sum(
            isinstance(
                worksheet.cell(
                    row_number,
                    column,
                ).value,
                str,
            )
            and worksheet.cell(
                row_number,
                column,
            ).value.strip() != ""
            and worksheet.cell(
                row_number,
                column,
            ).data_type != "f"
            for column in range(
                start_column,
                end_column + 1,
            )
        )

        if text_headers < 3:
            continue

        data_rows = []

        for candidate_row in range(
            row_number + 1,
            max_row + 1,
        ):
            populated = sum(
                worksheet.cell(
                    candidate_row,
                    column,
                ).value not in (None, "")
                for column in range(
                    start_column,
                    end_column + 1,
                )
            )

            if populated >= 2:
                data_rows.append(candidate_row)

            elif data_rows:
                break

        if len(data_rows) < 3:
            continue

        region = DataRegion(
            worksheet=worksheet.title,
            header_row=row_number,
            start_row=data_rows[0],
            end_row=data_rows[-1],
            start_column=start_column,
            end_column=end_column,
            headers=[
                str(value) if value is not None else ""
                for value in headers
            ],
        )

        regions.append(region)

        # V1 uses the first substantial table-like region per sheet.
        break

    return regions


from datetime import date, datetime


def classify_cell_value(value):
    """
    Classify populated Excel values into broad business data types.
    """
    if value in (None, ""):
        return None

    if isinstance(value, bool):
        return "boolean"

    if isinstance(value, (datetime, date)):
        return "date"

    if isinstance(value, (int, float)):
        return "numeric"

    if isinstance(value, str):
        return "text"

    return "other"


def normalise_record_value(value):
    """
    Normalise cell values for duplicate-record comparison.

    Blank values are represented consistently and text values
    have surrounding whitespace removed.
    """
    if value is None:
        return None

    if isinstance(value, str):
        return value.strip()

    return value